"""OXML helper smoke tests."""

from docx import Document

from app.services.renderer.oxml import add_page_number_fields, add_toc_field, set_page_border


def test_toc_and_page_fields_present_in_xml() -> None:
    doc = Document()
    p = doc.add_paragraph()
    add_toc_field(p)
    assert "TOC" in p._p.xml

    fp = doc.sections[0].footer.paragraphs[0]
    add_page_number_fields(fp)
    assert "PAGE" in fp._p.xml
    assert "NUMPAGES" in fp._p.xml

    set_page_border(doc.sections[0], "#2F5D8A")
    assert "pgBorders" in doc.sections[0]._sectPr.xml
