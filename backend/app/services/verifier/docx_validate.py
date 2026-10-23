"""Validate generated DOCX structure."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document


@dataclass
class DocxValidation:
    ok: bool
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    has_heading_styles: bool = False
    has_toc_field: bool = False
    has_page_fields: bool = False
    has_header: bool = False
    has_footer: bool = False


def check(path: Path | str) -> DocxValidation:
    path = Path(path)
    errors: list[str] = []
    warnings: list[str] = []
    try:
        doc = Document(str(path))
    except Exception as exc:
        return DocxValidation(ok=False, errors=[f"Cannot open DOCX: {exc}"])

    has_heading = False
    for p in doc.paragraphs:
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            has_heading = True
            break

    has_toc = False
    for p in doc.paragraphs:
        if "TOC" in p._p.xml and "instrText" in p._p.xml:
            has_toc = True
            break

    has_page = False
    has_header = False
    has_footer = False
    for section in doc.sections:
        header = section.header
        footer = section.footer
        if header.paragraphs and any(p.text.strip() for p in header.paragraphs):
            has_header = True
        if footer.paragraphs:
            has_footer = True
            for p in footer.paragraphs:
                xml = p._p.xml
                if "PAGE" in xml or "NUMPAGES" in xml:
                    has_page = True

    # Header row empty cells
    for table in doc.tables:
        if not table.rows:
            continue
        header = table.rows[0]
        for cell in header.cells:
            if not cell.text.strip():
                warnings.append("Empty table header cell found")
                break

    if not has_heading:
        errors.append("No Heading styles used")
    if not has_toc:
        warnings.append("TOC field not found")
    if not has_page:
        warnings.append("PAGE/NUMPAGES fields not found in footer")

    ok = len(errors) == 0
    return DocxValidation(
        ok=ok,
        errors=errors,
        warnings=warnings,
        has_heading_styles=has_heading,
        has_toc_field=has_toc,
        has_page_fields=has_page,
        has_header=has_header,
        has_footer=has_footer,
    )
