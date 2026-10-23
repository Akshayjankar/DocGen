"""Revision history — exactly 5 columns (Revision 2)."""

from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from app.models.document import DocumentMeta
from app.models.theme import Theme
from app.services.renderer.layout import revision_history_rows
from app.services.renderer.oxml import hex_to_word, set_cell_shading


def _rgb(hex_color: str) -> RGBColor:
    h = hex_to_word(hex_color)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_revision_history(
    doc: Document,
    theme: Theme,
    meta: DocumentMeta,
    *,
    page_break: bool = True,
) -> None:
    p = theme.palette
    doc.add_heading("Revision History", level=1)

    rows = revision_history_rows(meta, theme.layout)
    table = doc.add_table(rows=1 + len(rows), cols=5)
    table.style = "Table Grid"
    headers = ["#", "Version", "Date", "Author", "Reason for Change"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, p.primary)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    for r_i, values in enumerate(rows):
        for i, v in enumerate(values):
            cell = table.rows[r_i + 1].cells[i]
            cell.text = str(v).replace("\u00ad", "")
            set_cell_shading(cell, p.primary_tint)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()
    if page_break:
        doc.add_page_break()
