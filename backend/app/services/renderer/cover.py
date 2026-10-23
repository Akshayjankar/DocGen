"""Cover page with centered logo and filled-only release notice."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.models.document import DocumentMeta, RenderOptions
from app.models.theme import Theme
from app.services.renderer.dates import format_date
from app.services.renderer.layout import release_notice_rows
from app.services.renderer.oxml import hex_to_word, set_cell_shading


def _rgb(hex_color: str) -> RGBColor:
    h = hex_to_word(hex_color)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_cover(
    doc: Document,
    theme: Theme,
    meta: DocumentMeta,
    options: RenderOptions,
    logo_bytes: bytes | None = None,
    logo_path: str | Path | None = None,
) -> None:
    p = theme.palette

    if options.classification_label and meta.classification:
        ban = doc.add_paragraph()
        ban.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = ban.add_run(meta.classification.upper())
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = _rgb(p.muted)

    # Center logo only (corner not on cover unless show_corner_on_cover — DOCX: skip corner)
    cover_path = None
    if theme.logos and theme.logos.cover and theme.logos.cover.processed_path:
        cover_path = theme.logos.cover.processed_path
    if logo_bytes or logo_path or cover_path or theme.logo_path:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        try:
            if logo_bytes:
                import io

                run.add_picture(io.BytesIO(logo_bytes), width=Inches(2.2))
            else:
                path = logo_path or cover_path or theme.logo_path
                if path and Path(path).exists():
                    run.add_picture(str(path), width=Inches(2.2))
        except Exception:
            pass

    # Title as typed (uppercase only if layout.force_title_uppercase)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(28)
    display = meta.title
    if theme.layout and theme.layout.force_title_uppercase:
        display = meta.title.upper()
    tr = title.add_run(display)
    tr.font.name = theme.fonts.heading
    tr.font.size = Pt(30)
    tr.font.bold = True
    tr.font.color.rgb = _rgb(p.primary_dark)

    if meta.subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sub.add_run(meta.subtitle)
        sr.font.size = Pt(13)
        sr.font.color.rgb = _rgb(p.muted)

    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_s = format_date(meta.date or "", theme.layout)
    vr = ver.add_run(f"Version {meta.version or '1.0'} · {date_s}".strip(" ·"))
    vr.font.size = Pt(11)
    vr.font.color.rgb = _rgb(p.text)

    doc.add_paragraph()

    notice = doc.add_paragraph()
    nr = notice.add_run("Release Notice")
    nr.font.bold = True
    nr.font.size = Pt(12)
    nr.font.color.rgb = _rgb(p.primary_dark)

    rows = release_notice_rows(meta, theme.layout)
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        cell0 = table.rows[i].cells[0]
        cell1 = table.rows[i].cells[1]
        cell0.text = label
        cell1.text = value.replace("\u00ad", "")
        if i % 2 == 0:
            set_cell_shading(cell0, p.primary_tint)
            set_cell_shading(cell1, p.primary_tint)
        for cell in (cell0, cell1):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = theme.fonts.body
        for paragraph in cell0.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    doc.add_page_break()
