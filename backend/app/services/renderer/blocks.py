"""Map semantic Block[] → DOCX elements."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.models.blocks import Block
from app.models.theme import Theme
from app.services.renderer.inline_md import parse_inline, strip_inline_markers
from app.services.renderer.oxml import (
    add_seq_field,
    hex_to_word,
    set_cell_shading,
    set_left_cell_border,
)


def _rgb(hex_color: str) -> RGBColor:
    h = hex_to_word(hex_color)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_md_runs(
    para: Paragraph,
    text: str,
    *,
    theme: Theme,
    color: str,
    size_pt: float = 10.5,
    force_italic: bool = False,
) -> None:
    """Append runs with Markdown bold/italic/code resolved (markers omitted)."""
    for ir in parse_inline(text):
        run = para.add_run(ir.text)
        run.font.name = theme.fonts.mono if ir.code else theme.fonts.body
        run.font.size = Pt(size_pt)
        run.font.color.rgb = _rgb(color)
        run.font.bold = ir.bold
        run.font.italic = ir.italic or force_italic


def render_blocks(doc: Document, blocks: list[Block], source: str, theme: Theme) -> None:
    for block in blocks:
        _render_one(doc, block, source, theme)


def _render_one(doc: Document, block: Block, source: str, theme: Theme) -> None:
    p = theme.palette
    t = block.type

    if t == "heading":
        level = min(max(block.level or 1, 1), 3)
        text = strip_inline_markers(block.resolved_text(source))
        style_name = f"Docgen Heading {level}"
        try:
            para = doc.add_paragraph(text, style=style_name)
        except Exception:
            doc.add_heading(text, level=level)
        return

    if t == "paragraph":
        text = block.resolved_text(source)
        # Soft line breaks for single newlines inside a paragraph; blank lines already split blocks
        try:
            para = doc.add_paragraph(style="Docgen Body")
        except Exception:
            para = doc.add_paragraph()
        parts = text.replace("\r\n", "\n").split("\n")
        for i, part in enumerate(parts):
            _add_md_runs(para, part, theme=theme, color=p.text)
            if i < len(parts) - 1:
                para.add_run().add_break()
        para.paragraph_format.line_spacing = 1.45
        para.paragraph_format.space_after = Pt(6)
        return

    if t == "quote":
        text = block.resolved_text(source)
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(18)
        _add_md_runs(para, text, theme=theme, color=p.muted, force_italic=True)
        return

    if t == "divider":
        para = doc.add_paragraph("─" * 40)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    if t == "page_break":
        doc.add_page_break()
        return

    if t in ("steps", "bullets"):
        items = block.resolved_items(source)
        style = "List Number" if t == "steps" else "List Bullet"
        for item in items:
            para = doc.add_paragraph(style=style)
            _add_md_runs(para, item, theme=theme, color=p.text)
            para.paragraph_format.line_spacing = 1.45
        return

    if t == "table" and block.table:
        _render_table(doc, block, theme)
        return

    if t in ("callout_info", "callout_warning", "callout_success"):
        _render_callout(doc, block, source, theme)
        return

    if t == "code":
        _render_code(doc, block, source, theme)
        return

    if t == "figure":
        cap = doc.add_paragraph()
        cap.add_run("Figure ")
        add_seq_field(cap, "Figure")
        cap.add_run(f": {block.resolved_text(source) or block.label or ''}")
        try:
            cap.style = "Docgen Caption"
        except Exception:
            pass
        return

    # Fallback
    text = block.resolved_text(source)
    if text:
        doc.add_paragraph(text)


def _render_table(doc: Document, block: Block, theme: Theme) -> None:
    assert block.table is not None
    p = theme.palette
    headers = block.table.headers
    rows = block.table.rows
    cols = max(len(headers), max((len(r) for r in rows), default=0))
    if cols == 0:
        return
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True

    for i, h in enumerate(headers):
        if i >= cols:
            break
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, p.primary)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
                run.font.name = theme.fonts.body

    for r_i, row in enumerate(rows):
        for c_i in range(cols):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = row[c_i] if c_i < len(row) else ""
            if r_i % 2 == 1:
                set_cell_shading(cell, p.primary_tint)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.name = theme.fonts.body
    doc.add_paragraph()


def _render_callout(doc: Document, block: Block, source: str, theme: Theme) -> None:
    p = theme.palette
    color_map = {
        "callout_info": (p.info, p.primary_tint),
        "callout_warning": (p.warning, "#FFF8E7"),
        "callout_success": (p.success, "#E8F5E9"),
    }
    accent, bg = color_map[block.type]
    label = block.label or block.type.replace("callout_", "").upper()
    text = block.resolved_text(source)

    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, bg)
    set_left_cell_border(cell, accent, width_pt=18)
    cell.paragraphs[0].clear()
    label_run = cell.paragraphs[0].add_run(f"{label}  ")
    label_run.font.bold = True
    label_run.font.size = Pt(9)
    label_run.font.color.rgb = _rgb(accent)
    _add_md_runs(cell.paragraphs[0], text, theme=theme, color=p.text, size_pt=10)
    doc.add_paragraph()


def _render_code(doc: Document, block: Block, source: str, theme: Theme) -> None:
    p = theme.palette
    label = block.label or "CODE"
    text = block.resolved_text(source) or (block.text or "")

    label_p = doc.add_paragraph()
    lr = label_p.add_run(label)
    lr.font.name = theme.fonts.mono
    lr.font.size = Pt(8)
    lr.font.bold = True
    lr.font.color.rgb = _rgb(p.primary_dark)

    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "#F4F4F5")
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(text)
    run.font.name = theme.fonts.mono
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(p.text)
    doc.add_paragraph()
